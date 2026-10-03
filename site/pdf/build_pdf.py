#!/usr/bin/env python3
"""Typeset the paper as a Word document and PDF in the style of a Word paper.

Reads the built page (site/index.html), so the text cannot drift from the web
version, and lays it out the way the FIN11 note is set: Cambria, A4, a centred
title page with the abstract, a contents page, bold numbered headings, dash
lists, grey-ruled tables captioned above, figures captioned below, and a grey
page number at the foot of each page.

    python site/build.py          # the page this reads
    python site/pdf/figures.py    # the figure PNGs
    python site/pdf/build_pdf.py  # paper.docx, then paper.pdf via Word

The PDF step drives the installed Microsoft Word, which is also what fills in
the table of contents with page numbers.
"""
import re
import subprocess
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT,
                             WD_TAB_LEADER)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SITE = HERE.parent
PAGE = SITE / "index.html"
FIG = HERE / "fig"
DOCX = HERE / "paper.docx"
PDF = SITE / "paper.pdf"

TITLE = "Snowmaking and Day-Ahead Load Forecasts"
SUBTITLE = "A Pre-Registered Test in Four Markets"
AUTHOR = "Baptiste Cristofari"
AFFIL = ["Independent research", "August 2026"]
REPO = "https://github.com/baptistecristo/Electricity-demand-forecasting-in-Austria"

FONT = "Cambria"
MONO = "Consolas"
BLACK = RGBColor(0, 0, 0)
GRID = "BFBFBF"      # table rules
HEAD_FILL = "F2F2F2"  # table header row
HL_FILL = "DEEAF6"    # highlighted row


# ---------------------------------------------------------------- docx helpers

def set_font(run, name=FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def style_font(style, size, bold=False, name=FONT):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = BLACK
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]


def shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def table_borders(table, color=GRID):
    tblpr = table._element.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblpr.append(borders)


def cell_margins(table, top=40, bottom=40, left=90, right=90):
    tblpr = table._element.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblpr.append(mar)


def keep_with_next(par):
    par.paragraph_format.keep_with_next = True


def add_field(par, instr, size=None, color=None, name=FONT):
    """A Word field (PAGE, TOC...), left for Word to compute."""
    def run_with(el):
        r = par.add_run()
        set_font(r, name, size, color=color)
        r._element.append(el)
        return r
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    run_with(b); run_with(t); run_with(s)
    ph = par.add_run("1" if instr.strip() == "PAGE" else "")
    set_font(ph, name, size, color=color)
    run_with(e)


def add_hyperlink(par, url, text, bold=False, italic=False, size=None):
    rid = par.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), FONT)
    rpr.append(fonts)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    col = OxmlElement("w:color"); col.set(qn("w:val"), "1F3864"); rpr.append(col)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    if size:
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text
    r.append(t)
    link.append(r)
    par._p.append(link)


# ------------------------------------------------------------ html -> runs

def clean(text):
    return re.sub(r"\s+", " ", text)


def inline(par, node, bold=False, italic=False, mono=False, size=None):
    """Write an HTML fragment's inline content into a paragraph as runs."""
    for child in node.children:
        if isinstance(child, NavigableString):
            text = clean(str(child))
            if not text:
                continue
            r = par.add_run(text)
            set_font(r, MONO if mono else FONT,
                     (size or 11) * (0.9 if mono else 1) if (mono or size) else None,
                     bold=bold, italic=italic)
            continue
        if not isinstance(child, Tag):
            continue
        name = child.name
        if name in ("b", "strong"):
            inline(par, child, True, italic, mono, size)
        elif name in ("em", "i"):
            inline(par, child, bold, True, mono, size)
        elif name == "code":
            inline(par, child, bold, italic, True, size)
        elif name == "a":
            add_hyperlink(par, child.get("href", ""), clean(child.get_text()),
                          bold, italic, size)
        elif name == "br":
            par.add_run().add_break()
        else:
            inline(par, child, bold, italic, mono, size)


def strip_edges(par):
    """Trim the whitespace HTML leaves at the start and end of a paragraph."""
    runs = [r for r in par.runs if r.text]
    if runs:
        runs[0].text = runs[0].text.lstrip()
        # Only trim the end if a plain run really is last: a hyperlink after it
        # would otherwise lose the space that separates them.
        if par._p[-1] is runs[-1]._element:
            runs[-1].text = runs[-1].text.rstrip()


# ------------------------------------------------------------------ builder

class Paper:
    def __init__(self):
        self.doc = Document()
        self.fig_n = 0
        self._setup()

    def _setup(self):
        doc = self.doc
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Cm(2.5))
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)
        sec.footer_distance = Cm(1.25)

        st = doc.styles
        style_font(st["Normal"], 11)
        pf = st["Normal"].paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.08

        for name, size, before, after in (("Heading 1", 14, 24, 6),
                                          ("Heading 2", 12.5, 16, 4),
                                          ("Title", 18.5, 0, 0)):
            s = st[name]
            style_font(s, size, bold=True)
            s.paragraph_format.space_before = Pt(before)
            s.paragraph_format.space_after = Pt(after)
            s.paragraph_format.keep_with_next = True
            s.font.italic = False
            if name == "Title":
                # Word's Title carries a bottom border; the FIN11 title does not.
                ppr = s.element.get_or_add_pPr()
                for b in ppr.findall(qn("w:pBdr")):
                    ppr.remove(b)
        for lvl, (size, bold, indent) in {1: (11, True, 0), 2: (11, False, 0.6)}.items():
            try:
                s = st[f"toc {lvl}"]
            except KeyError:
                s = st.add_style(f"toc {lvl}", WD_STYLE_TYPE.PARAGRAPH)
                s.base_style = st["Normal"]
                s.element.set(qn("w:styleId"), f"TOC{lvl}")
            s.paragraph_format.tab_stops.add_tab_stop(
                Cm(16.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            style_font(s, size, bold=bold)
            s.paragraph_format.left_indent = Cm(indent)
            s.paragraph_format.space_after = Pt(2)

        # Grey page number, centred, on every page.
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_field(fp, "PAGE", size=9, color=RGBColor(0x80, 0x80, 0x80), name="Calibri")

    # -- blocks
    def para(self, node=None, text=None, align=None, size=None, after=None,
             bold=False, italic=False, style=None):
        p = self.doc.add_paragraph(style=style)
        if node is not None:
            inline(p, node, bold=bold, italic=italic, size=size)
            strip_edges(p)
        if text is not None:
            r = p.add_run(text)
            set_font(r, FONT, size, bold=bold, italic=italic)
        if align is not None:
            p.alignment = align
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        return p

    def heading(self, text, level):
        p = self.doc.add_heading(text, level=level)
        for r in p.runs:
            set_font(r, FONT, color=BLACK)
        return p

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def caption(self, node, label_re=r"^(Table|Figure) \d+\.$"):
        p = self.doc.add_paragraph()
        inline(p, node)
        strip_edges(p)
        p.paragraph_format.space_after = Pt(4)
        return p

    def dash_list(self, items, numbered=False):
        for i, li in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
            r = p.add_run(f"[{i}]\t" if numbered else "-\t")
            set_font(r, FONT)
            inline(p, li)
            strip_edges(p)
            if p.runs:
                p.runs[0].text = f"[{i}]\t" if numbered else "-\t"

    def code(self, node):
        lines = node.get_text().rstrip("\n").split("\n")
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_borders(tbl, "D9D9D9")
        cell_margins(tbl, 80, 80, 140, 140)
        cell = tbl.rows[0].cells[0]
        shade(cell, "F7F7F7")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for i, line in enumerate(lines):
            r = p.add_run(line)
            set_font(r, MONO, 9)
            if i < len(lines) - 1:
                r.add_break()
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)

    def table(self, node, caption=None):
        if caption is not None:
            keep_with_next(self.caption(caption))
        rows = node.find_all("tr")
        ncols = max(len(r.find_all(["th", "td"])) for r in rows)
        tbl = self.doc.add_table(rows=len(rows), cols=ncols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        table_borders(tbl)
        cell_margins(tbl)
        for i, tr in enumerate(rows):
            head = tr.parent.name == "thead"
            hl = "hl" in (tr.get("class") or [])
            for j, td in enumerate(tr.find_all(["th", "td"])):
                cell = tbl.rows[i].cells[j]
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                inline(p, td, bold=head, size=10)
                strip_edges(p)
                if head:
                    shade(cell, HEAD_FILL)
                elif hl:
                    shade(cell, HL_FILL)
            trpr = tbl.rows[i]._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if head:
                h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true")
                trpr.append(h)
            # keep a table on one page where it fits
            for cell in tbl.rows[i].cells:
                for p in cell.paragraphs:
                    p.paragraph_format.keep_with_next = i < len(rows) - 1
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)

    def figure(self, node):
        self.fig_n += 1
        img = FIG / f"fig{self.fig_n}.png"
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        keep_with_next(p)
        p.add_run().add_picture(str(img), width=Cm(15.5))
        cap = node.find("figcaption")
        if cap is not None:
            c = self.caption(cap)
            c.paragraph_format.space_after = Pt(12)

    # -- whole document
    def title_page(self, abstract):
        d = self.doc
        for _ in range(5):
            self.para(text="", after=12)
        t = self.para(text=TITLE, align=WD_ALIGN_PARAGRAPH.CENTER, size=18.5,
                      bold=True, after=2)
        self.para(text=SUBTITLE, align=WD_ALIGN_PARAGRAPH.CENTER, size=15,
                  bold=True, after=30)
        self.para(text=AUTHOR, align=WD_ALIGN_PARAGRAPH.CENTER, size=13, after=2)
        for line in AFFIL:
            self.para(text=line, align=WD_ALIGN_PARAGRAPH.CENTER, size=12, after=0)
        p = self.para(align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
        add_hyperlink(p, REPO, "Code and data on GitHub", size=12)
        for _ in range(3):
            self.para(text="", after=6)
        self.para(text="Abstract", align=WD_ALIGN_PARAGRAPH.CENTER, size=12,
                  bold=True, after=4)
        for par in abstract:
            lead = par.find("span", class_="lead")
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(6)
            if lead is not None:
                r = p.add_run(lead.get_text().strip() + ". ")
                set_font(r, FONT, 12, bold=True)
                lead.extract()
            inline(p, par, size=12)
            strip_edges(p)
            if lead is not None and p.runs:
                p.runs[0].text = p.runs[0].text.rstrip() + " "
        self.page_break()

    def contents(self):
        self.para(text="Contents", align=WD_ALIGN_PARAGRAPH.CENTER, size=15,
                  bold=True, after=18)
        p = self.doc.add_paragraph()
        add_field(p, 'TOC \\o "1-2" \\h \\z \\u')
        self.page_break()

    def body(self, container):
        blocks = []
        for el in container.children:
            if not isinstance(el, Tag):
                continue
            blocks.append(el)

        # Table captions sit above their table, as in the FIN11 note: pair each
        # caption with the nearest table before it.
        caption_for = {}
        last_table = None
        for el in blocks:
            if el.name == "div" and "twrap" in (el.get("class") or []):
                last_table = el
            elif el.name == "p" and "tcap" in (el.get("class") or []) and last_table is not None:
                caption_for[id(last_table)] = el
                last_table = None

        section_no = None
        for el in blocks:
            cls = el.get("class") or []
            if el.name in ("header",) or "tldr" in cls:
                continue
            if el.name == "span" and "chapter-label" in cls:
                section_no = int(el.get_text().split()[-1])
            elif el.name == "h2":
                title = clean(el.get_text()).strip()
                if section_no is not None:
                    title = f"{section_no}. {title}"
                section_no = None
                self.heading(title, 1)
            elif el.name == "h3":
                self.heading(clean(el.get_text()).strip(), 2)
            elif el.name == "p" and "tcap" in cls:
                continue  # placed with its table
            elif el.name == "p":
                self.para(el)
            elif el.name == "div" and "twrap" in cls:
                self.table(el.find("table"), caption_for.get(id(el)))
            elif el.name == "div" and "verdict" in cls:
                self.para(el)
            elif el.name == "div" and "foot" in cls:
                note, link = el.find_all("p")
                p = self.para(note)
                p.add_run(" ")
                a = link.find("a")
                add_hyperlink(p, a["href"], a.get_text())
            elif el.name == "figure":
                self.figure(el)
            elif el.name == "pre":
                self.code(el)
            elif el.name in ("ul", "ol"):
                self.dash_list(el.find_all("li", recursive=False),
                               numbered=el.name == "ol")
            else:
                raise SystemExit(f"unhandled block <{el.name} class={cls}>")


def main():
    soup = BeautifulSoup(PAGE.read_text(encoding="utf-8"), "html.parser")
    container = soup.select_one("main .container")
    abstract = container.select("section.tldr > p")

    paper = Paper()
    paper.title_page(abstract)
    paper.contents()
    paper.body(container)
    paper.doc.core_properties.title = f"{TITLE}: {SUBTITLE}"
    paper.doc.core_properties.author = AUTHOR
    paper.doc.save(DOCX)
    print(f"wrote {DOCX}")

    ps = f"""
$ErrorActionPreference = 'Stop'
$w = New-Object -ComObject Word.Application
$w.Visible = $false
try {{
  $d = $w.Documents.Open('{DOCX}')
  # Contents styled as in the FIN11 note: top level bold, sub-level indented.
  foreach ($i in @(-20, -21)) {{
    $st = $d.Styles.Item($i); $st.Font.Name = 'Cambria'; $st.Font.Size = 11
    $st.Font.Bold = ($i -eq -20); $st.ParagraphFormat.SpaceAfter = 3
  }}
  $d.Styles.Item(-21).ParagraphFormat.LeftIndent = 17
  # Columns sized to their content, then stretched to the text width.
  foreach ($t in $d.Tables) {{ $t.AllowAutoFit = $true; $t.AutoFitBehavior(1); $t.AutoFitBehavior(2) }}
  $d.TablesOfContents(1).Update()
  $d.Fields.Update() | Out-Null
  $d.Save()
  $d.ExportAsFixedFormat('{PDF}', 17, $false, 0, 0, 0, 0, 0, $true, $true, 1)
  $d.Close($false)
}} finally {{ $w.Quit() }}
"""
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    print(f"wrote {PDF}")


if __name__ == "__main__":
    main()
